"""Phase 3 helper: turn the blind coding CSVs into one friendly Excel file per coder, and import the filled files back.
  python src/03a2_coder_xlsx.py --build                         -> data_clean/coder_sheets/coder{1,2}.xlsx
  python src/03a2_coder_xlsx.py --import coder1 FILE.xlsx       -> data_clean/ai_coding_raw_coder1.csv (backup kept)
Coders must work independently and must NOT look up prices, multiples or premiums (rubric s0)."""
import argparse, shutil
from pathlib import Path
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
INPUT = ["product_test", "technical_asset_test", "revenue_test", "persistence_12m", "ai_intensity", "ai_subsector", "secondary_ai_subsector", "confidence",
         "rationale", "counterevidence_notes", "evidence_source_ids", "acquihire_ai_flag", "persistence_not_testable", "company_founded_year", "commercial_maturity"]
INSTR = [
 "AI CLASSIFICATION - BLIND CODING SHEET", "",
 "Work ALONE. Do not discuss with the other coder. Do NOT look up deal prices, multiples, premiums or how much the buyer paid.",
 "Classify each company AS IT WAS AT OR BEFORE the announcement date, using the description columns plus any pre-announcement sources (10-K, archived website, earlier press).",
 "", "FOUR TESTS (enter 1 = yes, 0 = no):",
 " product_test: would the main product lose its essential function if the ML/AI model were removed?",
 " technical_asset_test: documented proprietary trained models / ML infrastructure / dedicated ML research team?",
 " revenue_test: credible evidence AI/ML is a primary monetised product or driver of material revenue? (marketing claims alone = 0)",
 " persistence_12m: AI positioning documented at least 12 months before announcement? (enter 1/0, or set persistence_not_testable = 1)",
 "", "ai_intensity (REQUIRED): 0 = Non-AI | 1 = AI-enabled (AI improves product but is not fundamental) | 2 = AI-native (removing AI changes what is sold) | 3 = Frontier / core AI infrastructure",
 "ai_subsector (REQUIRED only if intensity is 2 or 3): 1 Foundation/GenAI, 2 AI dev tools/MLOps/data infra, 3 Enterprise AI software, 4 Vertical AI, 5 Vision/NLP/Conversational, 6 Autonomous/Robotics, 7 AI cybersecurity, 8 AI semis/accelerators, 9 Traditional ML",
 "confidence: High / Medium / Low.   rationale: one or two sentences (required).   acquihire_ai_flag: 1 only if product shut down / founders-only / IP licence.",
 "When unsure between two levels pick the lower one, set confidence = Low, and explain in rationale. Enterprise software + an AI assistant = Level 1. Cybersecurity with ML detection = Level 1 unless the product IS a behavioural model.",
 "Full rubric: docs/ai_rubric.md. Fill ONLY the yellow columns. Do not edit, sort-delete or add rows.",
]


def build():
    out = ROOT / "data_clean" / "coder_sheets"; out.mkdir(parents=True, exist_ok=True)
    for n in (1, 2):
        d = pd.read_csv(ROOT / "data_clean" / f"ai_coding_raw_coder{n}.csv", dtype=str, keep_default_na=False)
        wb = Workbook(); ws = wb.active; ws.title = "Instructions"
        for i, t in enumerate(INSTR, 1):
            ws.cell(i, 1, t).alignment = Alignment(wrap_text=True)
        ws["A1"].font = Font(bold=True, size=14); ws.column_dimensions["A"].width = 150
        c = wb.create_sheet("Coding"); c.append(list(d.columns))
        for r in d.itertuples(index=False): c.append(list(r))
        yellow = PatternFill("solid", fgColor="FFF2CC"); grey = PatternFill("solid", fgColor="D9D9D9")
        for j, col in enumerate(d.columns, 1):
            c.cell(1, j).font = Font(bold=True); c.cell(1, j).fill = yellow if col in INPUT else grey
            c.column_dimensions[c.cell(1, j).column_letter].width = 60 if "description" in col or "snapshot" in col else 16
            for i in range(2, len(d) + 2): c.cell(i, j).alignment = Alignment(wrap_text=True, vertical="top")
        idx = {col: j for j, col in enumerate(d.columns, 1)}
        def dv(col, formula):
            v = DataValidation(type="list", formula1=formula, allow_blank=True); c.add_data_validation(v)
            L = c.cell(1, idx[col]).column_letter; v.add(f"{L}2:{L}{len(d)+1}")
        for col in ("product_test", "technical_asset_test", "revenue_test", "persistence_12m", "acquihire_ai_flag", "persistence_not_testable"): dv(col, '"0,1"')
        dv("ai_intensity", '"0,1,2,3"'); dv("ai_subsector", '"1,2,3,4,5,6,7,8,9"'); dv("secondary_ai_subsector", '"1,2,3,4,5,6,7,8,9"'); dv("confidence", '"High,Medium,Low"')
        c.freeze_panes = "D2"; wb.save(out / f"coder{n}.xlsx"); print("wrote", out / f"coder{n}.xlsx", len(d), "rows")


def imp(who, f):
    d = pd.read_excel(f, sheet_name="Coding", dtype=str).fillna("")
    d = d.apply(lambda s: s.str.replace(r"\.0$", "", regex=True) if s.name in INPUT[:8] + INPUT[11:13] + ["company_id", "coder_row"] else s)
    tgt = ROOT / "data_clean" / f"ai_coding_raw_{who}.csv"
    orig = pd.read_csv(tgt, dtype=str, keep_default_na=False)
    assert list(d.columns) == list(orig.columns) and len(d) == len(orig), "sheet structure changed"
    assert (d["company_id"].tolist() == orig["company_id"].tolist()), "rows reordered or edited"
    shutil.copy(tgt, tgt.with_suffix(".backup.csv")); d.to_csv(tgt, index=False); print("imported", who, int((d["ai_intensity"] != "").sum()), "coded rows ->", tgt)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--build", action="store_true"); ap.add_argument("--import", dest="imp", nargs=2, metavar=("coder1|coder2", "FILE"))
    a = ap.parse_args()
    if a.build: build()
    elif a.imp: imp(*a.imp)
